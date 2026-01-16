import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill

def compare_and_color_excel_by_index(
    ref_file_path: str,
    comp_file_path: str,
    ref_cols_indices: list,
    comp_cols_indices: list,
    output_file_path: str,
    comp_sheet_to_process: str = None
):
    """
    Compara uma coluna de um arquivo Excel com uma coluna de outro (usando o primeiro índice de cada lista).
    Para cada correspondência, anota a planilha de origem e o valor.
    Gera um novo arquivo com as linhas coloridas e as novas informações.

    Args:
        ref_file_path (str): Caminho para o arquivo Excel de referência.
        comp_file_path (str): Caminho para o arquivo Excel a ser comparado.
        ref_cols_indices (list): Lista com o índice (1-based) da coluna do arquivo de referência.
        comp_cols_indices (list): Lista com o índice (1-based) da coluna do arquivo de comparação.
        output_file_path (str): Caminho para salvar o arquivo Excel de saída.
        comp_sheet_to_process (str, optional): Nome da planilha a ser processada. Se None, processa todas.
    """
    print("Iniciando a comparação de coluna única...")

    # --- Etapa 1: Mapear todos os dados de referência para sua planilha de origem ---
    print(f"Lendo e mapeando o arquivo de referência: {ref_file_path}")
    try:
        ref_excel = pd.ExcelFile(ref_file_path)
    except FileNotFoundError:
        print(f"Erro: Arquivo de referência não encontrado em '{ref_file_path}'")
        return

    reference_data_map = {}
    ref_col_0based = ref_cols_indices[0] - 1

    for sheet_name in ref_excel.sheet_names:
        print(f"  - Processando planilha de referência: {sheet_name}")
        try:
            df_ref = pd.read_excel(ref_excel, sheet_name=sheet_name, header=None)
            for _, row in df_ref.iterrows():
                value = row.iloc[ref_col_0based]
                
                if pd.notna(value):
                    key = str(value)
                    if key not in reference_data_map:
                        reference_data_map[key] = sheet_name
        except IndexError:
            print(f"Aviso: Pulando planilha '{sheet_name}' pois a coluna {ref_cols_indices[0]} não foi encontrada.")
            continue
    
    print(f"Total de {len(reference_data_map)} valores de referência únicos mapeados.")

    # --- Etapa 2: Carregar o arquivo de comparação e processá-lo ---
    print(f"\nLendo e processando arquivo de comparação: {comp_file_path}")
    try:
        comp_workbook = openpyxl.load_workbook(comp_file_path)
    except FileNotFoundError:
        print(f"Erro: Arquivo de comparação não encontrado em '{comp_file_path}'")
        return

    sheets_to_process = []
    if comp_sheet_to_process:
        if comp_sheet_to_process in comp_workbook.sheetnames:
            sheets_to_process.append(comp_sheet_to_process)
        else:
            print(f"Erro: Planilha '{comp_sheet_to_process}' não encontrada no arquivo '{comp_file_path}'.")
            print(f"Planilhas disponíveis: {comp_workbook.sheetnames}")
            return
    else:
        sheets_to_process = comp_workbook.sheetnames

    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    gray_fill = PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid")

    for sheet_name in sheets_to_process:
        worksheet = comp_workbook[sheet_name]
        print(f"  - Processando planilha de comparação: '{sheet_name}'")

        header_col_1 = worksheet.max_column + 1
        header_col_2 = worksheet.max_column + 2
        worksheet.cell(row=1, column=header_col_1).value = "Planilha de Origem (Referência)"
        worksheet.cell(row=1, column=header_col_2).value = "Valor Encontrado"

        comp_col_idx = comp_cols_indices[0]

        if comp_col_idx > worksheet.max_column:
            print(f"Aviso: Pulando planilha '{sheet_name}' pois a coluna {comp_col_idx} não foi encontrada.")
            continue

        for row_idx in range(2, worksheet.max_row + 1):
            # Print para feedback em tempo real da linha sendo processada
            print(f"\r    -> Verificando linha {row_idx-1} de {worksheet.max_row-1}...", end="", flush=True)

            cell_value = worksheet.cell(row=row_idx, column=comp_col_idx).value
            
            fill_style = None
            if pd.notna(cell_value):
                comparison_key = str(cell_value)
                if comparison_key in reference_data_map:
                    fill_style = green_fill
                    found_in_sheet = reference_data_map[comparison_key]
                    worksheet.cell(row=row_idx, column=header_col_1).value = found_in_sheet
                    worksheet.cell(row=row_idx, column=header_col_2).value = comparison_key
                else:
                    fill_style = gray_fill

            if fill_style:
                for col_idx in range(1, header_col_1):
                    worksheet.cell(row=row_idx, column=col_idx).fill = fill_style
        
        print() # Adiciona uma nova linha após o término do loop de progresso

    # --- Etapa 3: Salvar o arquivo resultante ---
    try:
        comp_workbook.save(output_file_path)
        print(f"\nComparação concluída! Arquivo salvo em: '{output_file_path}'")
    except Exception as e:
        print(f"Erro ao salvar o arquivo: {e}")

if __name__ == '__main__':
    # --- CONFIGURE AQUI ---
    ARQUIVO_REFERENCIA = "C:\\Users\\walyson.ferreira\\Desktop\\Arquivos-SINAPI-SICRO-ORSE\\agrupado\\agrupado_20260107_100254.xlsx"
    ARQUIVO_COMPARACAO = "C:\\Users\\walyson.ferreira\\Downloads\\Planilha_Orçamentária.xlsx"
    # ARQUIVO_REFERENCIA = "C:\\Users\\walyson.ferreira\\Downloads\\Planilha_Orçamentária.xlsx"
    # ARQUIVO_COMPARACAO = "C:\\Users\\walyson.ferreira\\Desktop\\Arquivos-SINAPI-SICRO-ORSE\\agrupado\\agrupado_20260107_100254.xlsx"
    
    
    PLANILHA_ALVO_COMPARACAO = 'Curva ABC'
    
    COLUNAS_REFERENCIA = [1] # Coluna A
    COLUNAS_COMPARACAO = [2] # Coluna B
    
    
    ARQUIVO_SAIDA = "resultado_comparacao.xlsx"
    
    # --- FIM DA CONFIGURAÇÃO ---

    compare_and_color_excel_by_index(
        ref_file_path=ARQUIVO_REFERENCIA,
        comp_file_path=ARQUIVO_COMPARACAO,
        ref_cols_indices=COLUNAS_REFERENCIA,
        comp_cols_indices=COLUNAS_COMPARACAO,
        output_file_path=ARQUIVO_SAIDA,
        comp_sheet_to_process=PLANILHA_ALVO_COMPARACAO
    )