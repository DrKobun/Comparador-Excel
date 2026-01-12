import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill

class ABCCurveApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Calculadora de Curva ABC")
        self.root.geometry("500x500")

        self.file_path = None
        self.excel_file = None
        self.df = None

        # --- Widgets ---
        self.main_frame = tk.Frame(self.root, padx=10, pady=10)
        self.main_frame.pack(fill=tk.BOTH, expand=True)

        # File Selection
        self.file_button = tk.Button(self.main_frame, text="Selecionar Arquivo Excel", command=self.load_excel_file)
        self.file_button.pack(pady=10)

        self.file_label = tk.Label(self.main_frame, text="Nenhum arquivo selecionado", wraplength=380)
        self.file_label.pack()

        # Sheet Selection
        self.sheet_frame = tk.LabelFrame(self.main_frame, text="Seleção de Planilha", padx=10, pady=10)
        self.sheet_frame.pack(pady=10, fill=tk.X)

        self.sheet_label = tk.Label(self.sheet_frame, text="Planilha:")
        self.sheet_label.grid(row=0, column=0, sticky=tk.W, pady=5)
        self.sheet_var = tk.StringVar()
        self.sheet_dropdown = ttk.Combobox(self.sheet_frame, textvariable=self.sheet_var, state="disabled")
        self.sheet_dropdown.grid(row=0, column=1, sticky=tk.EW)
        self.sheet_frame.grid_columnconfigure(1, weight=1)
        self.sheet_dropdown.bind("<<ComboboxSelected>>", self.load_sheet_data)


        # Column Selection
        self.column_frame = tk.LabelFrame(self.main_frame, text="Selecionar Colunas", padx=10, pady=10)
        self.column_frame.pack(pady=10, fill=tk.X)

        self.desc_label = tk.Label(self.column_frame, text="Descrição:")
        self.desc_label.grid(row=0, column=0, sticky=tk.W, pady=5)
        self.desc_var = tk.StringVar()
        self.desc_dropdown = ttk.Combobox(self.column_frame, textvariable=self.desc_var, state="disabled")
        self.desc_dropdown.grid(row=0, column=1, sticky=tk.EW)

        self.qty_label = tk.Label(self.column_frame, text="Quantidade:")
        self.qty_label.grid(row=1, column=0, sticky=tk.W, pady=5)
        self.qty_var = tk.StringVar()
        self.qty_dropdown = ttk.Combobox(self.column_frame, textvariable=self.qty_var, state="disabled")
        self.qty_dropdown.grid(row=1, column=1, sticky=tk.EW)

        self.value_label = tk.Label(self.column_frame, text="Valor Unitário:")
        self.value_label.grid(row=2, column=0, sticky=tk.W, pady=5)
        self.value_var = tk.StringVar()
        self.value_dropdown = ttk.Combobox(self.column_frame, textvariable=self.value_var, state="disabled")
        self.value_dropdown.grid(row=2, column=1, sticky=tk.EW)
        
        self.column_frame.grid_columnconfigure(1, weight=1)

        # Calculate Button
        self.calculate_button = tk.Button(self.main_frame, text="Calcular Curva ABC", command=self.calculate_abc, state="disabled")
        self.calculate_button.pack(pady=20)

        self.status_var = tk.StringVar()
        self.status_bar = ttk.Label(self.root, textvariable=self.status_var, relief=tk.SUNKEN, anchor=tk.W)
        self.status_bar.pack(side=tk.BOTTOM, fill=tk.X)

    def load_excel_file(self):
        file_path_selected = filedialog.askopenfilename(
            title="Selecione o arquivo Excel com os dados dos produtos",
            filetypes=[("Excel files", "*.xlsx *.xls *.xlsm")]
        )
        if file_path_selected:
            try:
                self.clear_selections()
                self.file_path = file_path_selected
                self.excel_file = pd.ExcelFile(self.file_path)
                sheet_names = self.excel_file.sheet_names
                self.sheet_dropdown['values'] = sheet_names
                self.sheet_dropdown.config(state="readonly")
                self.file_label.config(text=self.file_path.split('/')[-1])
                
                if len(sheet_names) == 1:
                    self.sheet_var.set(sheet_names[0])
                    self.load_sheet_data()

            except Exception as e:
                messagebox.showerror("Erro", f"Erro ao ler o arquivo Excel:\n{e}")
                self.clear_selections()


    def load_sheet_data(self, event=None):
        sheet_name = self.sheet_var.get()
        if sheet_name and self.excel_file:
            try:
                self.df = pd.read_excel(self.excel_file, sheet_name=sheet_name)
                self.populate_column_dropdowns()
                self.calculate_button.config(state="normal")
            except Exception as e:
                messagebox.showerror("Erro", f"Erro ao carregar a planilha '{sheet_name}':\n{e}")
                self.clear_selections()


    def populate_column_dropdowns(self):
        self.desc_var.set('')
        self.qty_var.set('')
        self.value_var.set('')
        columns = list(self.df.columns)
        self.desc_dropdown['values'] = columns
        self.qty_dropdown['values'] = columns
        self.value_dropdown['values'] = columns
        self.desc_dropdown.config(state="readonly")
        self.qty_dropdown.config(state="readonly")
        self.value_dropdown.config(state="readonly")

    def clear_selections(self):
        self.file_path = None
        self.excel_file = None
        self.df = None
        
        self.file_label.config(text="Nenhum arquivo selecionado")
        
        self.sheet_var.set('')
        self.sheet_dropdown['values'] = []
        self.sheet_dropdown.config(state="disabled")

        self.desc_var.set('')
        self.qty_var.set('')
        self.value_var.set('')
        self.desc_dropdown['values'] = []
        self.qty_dropdown['values'] = []
        self.value_dropdown['values'] = []
        self.desc_dropdown.config(state="disabled")
        self.qty_dropdown.config(state="disabled")
        self.value_dropdown.config(state="disabled")

        self.calculate_button.config(state="disabled")

    def calculate_abc(self):
        sheet_name = self.sheet_var.get()
        desc_col = self.desc_var.get()
        qty_col = self.qty_var.get()
        value_col = self.value_var.get()

        if not sheet_name:
            messagebox.showerror("Erro", "Por favor, selecione uma planilha.")
            return
            
        if not (desc_col and qty_col and value_col):
            messagebox.showerror("Erro", "Por favor, selecione todas as colunas.")
            return

        try:
            self.status_var.set("Iniciando cálculo...")
            self.root.update_idletasks()

            # Ensure numeric columns are treated as such
            self.status_var.set("Convertendo colunas para numérico...")
            self.root.update_idletasks()
            self.df[qty_col] = pd.to_numeric(self.df[qty_col])
            self.df[value_col] = pd.to_numeric(self.df[value_col])

            # Calculate Total Value
            self.status_var.set("Calculando valor total...")
            self.root.update_idletasks()
            self.df['Valor Total'] = self.df[qty_col] * self.df[value_col]

            # Sort by Total Value
            self.status_var.set("Ordenando produtos...")
            self.root.update_idletasks()
            df_sorted = self.df.sort_values(by='Valor Total', ascending=False)

            # Calculate Cumulative Percentage
            self.status_var.set("Calculando percentual acumulado...")
            self.root.update_idletasks()
            total_sum = df_sorted['Valor Total'].sum()
            df_sorted['Percentual Acumulado'] = (df_sorted['Valor Total'].cumsum() / total_sum) * 100

            # Classify ABC
            self.status_var.set("Classificando itens em A, B e C...")
            self.root.update_idletasks()
            def classify_abc(perc):
                if perc <= 80:
                    return 'A'
                elif 80 < perc <= 95:
                    return 'B'
                else:
                    return 'C'

            df_sorted['Categoria'] = df_sorted['Percentual Acumulado'].apply(classify_abc)

            # Prepare final DataFrame
            output_df = df_sorted[[
                desc_col,
                qty_col,
                'Valor Total',
                'Categoria',
                'Percentual Acumulado'
            ]]
            output_df.rename(columns={desc_col: 'Descrição', qty_col: 'Quantidade'}, inplace=True)

            self.status_var.set("Aguardando local para salvar o arquivo...")
            self.root.update_idletasks()
            # Save to new Excel file
            output_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                title="Salvar Arquivo de Curva ABC"
            )

            if not output_path:
                self.status_var.set("Operação cancelada.")
                return # User cancelled save dialog

            self.status_var.set("Salvando arquivo Excel...")
            self.root.update_idletasks()
            output_df.to_excel(output_path, index=False, sheet_name='Curva ABC')

            # Apply formatting
            self.status_var.set("Aplicando formatação...")
            self.root.update_idletasks()
            workbook = openpyxl.load_workbook(output_path)
            sheet = workbook['Curva ABC']
            
            green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

            # Find the 'Categoria' column index
            header = [cell.value for cell in sheet[1]]
            try:
                category_col_index = header.index('Categoria') + 1
            except ValueError:
                messagebox.showerror("Erro", "Não foi possível encontrar a coluna 'Categoria' no arquivo de saída.")
                self.status_var.set("Erro ao formatar.")
                return

            for row_idx, row in enumerate(sheet.iter_rows(min_row=2, max_row=sheet.max_row), start=2):
                category_cell = sheet.cell(row=row_idx, column=category_col_index)
                if category_cell.value == 'A':
                    for cell in row:
                        cell.fill = green_fill
            
            # Adjust column widths
            for col in sheet.columns:
                max_length = 0
                column = col[0].column_letter # Get the column name
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = (max_length + 2)
                sheet.column_dimensions[column].width = adjusted_width


            workbook.save(output_path)
            self.status_var.set("Pronto.")
            messagebox.showinfo("Sucesso", f"Arquivo de Curva ABC salvo em:\n{output_path}")

        except KeyError as e:
            messagebox.showerror("Erro de Coluna", f"A coluna selecionada não foi encontrada: {e}")
            self.status_var.set("Erro.")
        except Exception as e:
            messagebox.showerror("Erro", f"Ocorreu um erro durante o cálculo:\n{e}")
            self.status_var.set("Erro.")
        finally:
            self.root.after(5000, lambda: self.status_var.set("")) # Clear status bar after 5 seconds



if __name__ == "__main__":
    root = tk.Tk()
    app = ABCCurveApp(root)
    root.mainloop()
